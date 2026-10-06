"""Rebuild one district's living income benchmark workbook (v0.02) from the survey + market data.
Usage: python build_benchmark.py <district> <src.xlsx> <participants.csv> <market_tidy_v1_1.csv> <out.xlsx>
Keeps all of the tool's formulas; only inputs are replaced and a few clear formula bugs fixed (listed in CHANGES)."""
import sys, copy, numpy as np, pandas as pd, openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
district, src, part, mk, out = sys.argv[1:6]
wb = openpyxl.load_workbook(src)
CH = []
P = pd.read_csv(part)
# ---- 1. participants -> CLS - Input data (rows) and Decent housing (rows)
cls = wb["CLS - Input data"]; dh = wb["Decent housing"]
rowmap = {16:"Age",17:"Gender",18:"HHsize",19:"Adults",20:"FT",21:"PT",22:"Children",23:"Rooms",24:"Area_m2",
          26:"Rent",27:"Water",28:"Electricity",29:"CookingFuel",31:"Healthcare",32:"Education",33:"Transport"}
for j, (_, r) in enumerate(P.iterrows()):
    for row, col in rowmap.items():
        v = r[col]; v = None if pd.isna(v) else (round(float(v), 2) if not isinstance(v, str) else v)
        cls.cell(row, 4 + j).value = v
    for row, col in {27:"d_include",28:"d_walls",29:"d_roof",30:"d_floor",31:"d_toilet",32:"d_water",33:"d_cooking",34:"d_electricity"}.items():
        dh.cell(row, 5 + j).value = r[col]
CH.append(f"Participants: replaced the 15 CLS participants with the 15 most complete {district} households (ids {list(P['_id'])})")
# ---- 2. boundaries / reference household
sb = wb["1. Setting the boundaries"]
sb["D8"] = "Rural"; sb["D25"] = 1; sb["D26"] = 1; sb["D27"] = 3; sb["D28"] = 2
CH.append("Zone set to Rural; personalised reference household = 1 adult male + 1 adult female + 3 children, 2 full-time workers")
out_ws = wb["OUTPUT_Living wage"]
if isinstance(out_ws["E7"].value, str) and out_ws["E7"].value.startswith("="): pass
else: out_ws["E7"] = "Rural"
# ---- 3. food prices from the market survey
M = pd.read_csv(mk); M = M[M.district == district]
M = M[M.outlier_flag.isna()]   # drop prices flagged as outliers (the tool prices at the LOWEST observed price)
def prices(item, unit):
    x = M[(M["item"] == item) & M.price_ugx.notna()]
    std = x[(x.price_std_unit == unit) & x.price_std.notna()].price_std.tolist()
    if std: return sorted(std), "weighed/standard"
    g = x[(x.local_factor_unit == unit) & x.price_std_guess.notna()].price_std_guess.tolist()
    if g:
        med = float(np.median(g)); g = [v for v in g if v >= 0.5 * med]   # guessed conversions: drop values below half their median (tool uses the LOWEST price)
        return sorted(g), "guessed local unit (values < 50% of median dropped)"
    return [], "none"
MAP = {"Rice 1":("rice","kg"),"Rice 2":("rice","kg"),"Rice 3":("rice","kg"),"Rice 4":("rice","kg"),"Bread":("bread","kg"),
 "chapati":("naan","kg"),"Cassava":("cassava","kg"),"Potatoes":("irish_potatoes","kg"),"Yams":("yams","kg"),"Pumpkin":("pumpkin","kg"),
 "Matooke":("banana","kg"),"Groundnuts":("ground_nuts","kg"),"fresh Beans":("fresh_beans","kg"),"Peas":("cow_peas","kg"),
 "Soy":("dried_soybeans","kg"),"drybeans":("dried_beans","kg"),"Milk":("milk","litre"),"Youghurt":("yoghurt","litre"),
 "Butter":("margarine","kg"),"Chicken":("chicken","kg"),"beef":("beaf","kg"),"Goat/Chevon":("goat","kg"),"Ngege Fresh":("fish","kg"),
 "Nakati":("nakati","kg"),"Doodo":("doodo","kg"),"Bbuga":("ebugga","kg"),"Cabbage":("cabbage","kg"),"Water Melon":("watermelon","kg"),
 "Mangoes":("mangoes","kg"),"Bananas":("sweet_banana","kg"),"Cooking Oil":("cookingoil","litre")}
fi = wb["Food costs - Input data"]; log = []
for r in range(11, 64):
    name = fi.cell(r, 3).value
    if name in MAP:
        vals, basis = prices(*MAP[name])
        if vals:
            for c in range(4, 29): fi.cell(r, c).value = None
            for k, v in enumerate(vals[:25]): fi.cell(r, 4 + k).value = round(float(v), 0)
            log.append((name, MAP[name][0], len(vals), min(vals), np.median(vals), basis))
        else: log.append((name, MAP[name][0], 0, None, None, "no usable market price - tool value kept"))
    elif name == "Eggs":
        x = M[(M["item"] == "eggs") & M.price_ugx.notna()]; t = x[x.price_std_unit == "per tray"].price_std.tolist()
        if t:
            kg = [v / 1.8 for v in sorted(t)]       # assumption: 30 eggs x ~60 g = 1.8 kg per tray
            for c in range(4, 29): fi.cell(r, c).value = None
            for k, v in enumerate(kg[:25]): fi.cell(r, 4 + k).value = round(v, 0)
            log.append((name, "eggs", len(kg), min(kg), np.median(kg), "per tray / 1.8 kg (assumed)"))
    else: log.append((name, "-", 0, None, None, "not in market survey - tool value kept"))
pd.DataFrame(log, columns=["tool_item","market_item","n_prices","lowest","median","basis"]).to_csv(out.replace(".xlsx","_food_price_log.csv"), index=False)
CH.append("Food prices: replaced by July 2025 district market prices for items with usable prices (see food price log); other items keep the tool's original values")
# ---- 4. clear formula fixes
hc = wb["2. Housing costs"]
for rr in range(30, 60):
    v = str(hc.cell(rr, 7).value)
    if "AVERAGE(" in v and ",1)" in v:
        hc.cell(rr, 7).value = v.replace(",1)", ")", 1); CH.append(f"2. Housing costs G{rr}: removed the stray ',1' from the average of cost per room")
# heights (Nakaseke workbook had 2.2 m / 2.0 m; decision 2026-10-05: use Mukono's values)
ni = wb["Nutritional input"]
if ni["E12"].value != 1.6:
    CH.append(f"Nutritional input: average heights set from {ni['E12'].value}/{ni['G12'].value} m to 1.6 m (men) / 1.5 m (women)"); ni["E12"] = 1.6; ni["G12"] = 1.5
o = out_ws
o["G31"] = '=IFERROR(E31/$E$10,"")'; CH.append("OUTPUT G31: USD margins now = E31/E10 (was pointing to the elder-care cell F22)")
o["C48"] = "Cost of elder care"; o["C49"] = "Margins (unexpected events)"
for col, src_cell in (("E", "E48"), ("G", "G48")):
    o[f"{col}49"] = copy.copy(o[src_cell]._style) and None
o["E49"] = '=IFERROR(E31*12,"")'; o["G49"] = '=IFERROR(G31*12,"")'
for c in "CDEFGHIJ": o[f"{c}49"]._style = copy.copy(o[f"{c}48"]._style)
o["E50"] = "=SUM(E45:E49)"; o["G50"] = "=SUM(G45:G49)"
CH.append("OUTPUT yearly table: relabelled row 48 as elder care, added the missing margins row 49, totals now include it")
# ---- 5. restore dropdowns that openpyxl drops
dv = DataValidation(type="list", formula1="='Background calculations'!$B$10:$B$11", allow_blank=True); o.add_data_validation(dv); dv.add("E15")
dv2 = DataValidation(type="list", formula1="='Background calculations'!$B$202:$B$206", allow_blank=True); wb["Nutritional input"].add_data_validation(dv2); dv2.add("E112:E130")
wb.calculation.fullCalcOnLoad = True
wb.save(out); open(out.replace(".xlsx","_changes.txt"), "w").write("\n".join(CH))
print("\n".join(CH))
