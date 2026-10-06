"""moreFigures - group food_prices_gap: redo of the old notebooks marketPrices, modalDietCosting, FoodAnalysis,
foodAnslysisNew, NFNH and incomeGap-Mukono on the cleaned data and the v0.02 living income benchmark.

Run from the repository root:   python scripts/morefigures/food_prices_gap.py
Writes to paper1/figures/moreFigures/<old folder>/ :
  Food_Analysis, costing, food_cleaning, nfnh_identification, food_constraints, market_prices,
  gap_drivers_mukono, living_income_gap_mukono, living_income_gap_mukono_outliers
Model diet / NFNH tables: the v0.02 benchmark workbooks are not in the repo. They are rebuilt here in a temporary
folder with the repo's own scripts/benchmark/ (same inputs, same rules) and recalculated with LibreOffice; the
rebuilt totals are checked against the published values in docs/logs/benchmark-rebuild-log-*.md. Without
LibreOffice (or if the check fails) only the published component values from the logs are written.
"""
import os, re, sys, textwrap, shutil, subprocess, tempfile, warnings
from pathlib import Path
import numpy as np, pandas as pd
from scipy import stats

warnings.filterwarnings("ignore")
R = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(R / "scripts" / "obj1"))
from li_style import *                                   # noqa: E402,F401
from matplotlib.colors import LinearSegmentedColormap, TwoSlopeNorm   # noqa: E402

DER, RAW, OUT = R / "data" / "derived", R / "data" / "raw", R / "paper1" / "figures" / "moreFigures"
BENCH = R / "scripts" / "benchmark"
DISTS = ["mukono", "nakaseke"]
num = lambda x: pd.to_numeric(x, errors="coerce")


def odir(name):
    p = OUT / name; p.mkdir(parents=True, exist_ok=True); return p


def to_tex(df, path, caption, label, fmt=None):
    """Minimal booktabs LaTeX table (no jinja2 dependency)."""
    fmt = fmt or {}
    esc = lambda s: str(s).replace("&", r"\&").replace("%", r"\%").replace("_", r"\_").replace("#", r"\#").replace(">=", r"$\geq$").replace("<", r"$<$")
    def cell(c, v):
        if pd.isna(v): return "--"
        if c in fmt: return fmt[c].format(v)
        if isinstance(v, (float, np.floating)): return f"{v:,.0f}" if abs(v) >= 100 else f"{v:.1f}"
        return esc(v)
    lines = [r"\begin{table}[htbp]", r"\centering", rf"\caption{{{esc(caption)}}}", rf"\label{{{label}}}",
             r"\begin{tabular}{" + "l" + "r" * (df.shape[1] - 1) + "}", r"\toprule",
             " & ".join(esc(c) for c in df.columns) + r" \\", r"\midrule"]
    for _, r in df.iterrows():
        lines.append(" & ".join(cell(c, r[c]) for c in df.columns) + r" \\")
    lines += [r"\bottomrule", r"\end{tabular}", r"\end{table}"]
    Path(path).write_text("\n".join(lines) + "\n")


# ------------------------------------------------------------------ data
h = pd.read_csv(DER / "Household_Analysis_v1.csv")
sv = pd.read_csv(DER / "Survey_Cleaned_v1.csv", low_memory=False)
sv = sv[sv.District.isin(DISTS)].merge(h[["_id"]], on="_id")
assert len(h) == 597 and len(sv) == 597
NHH = h.district.value_counts().to_dict()
mk = pd.read_csv(DER / "Market_Prices_Tidy_v1_1.csv")

# ================================================================== 1. benchmark v0.02 (rebuild + check)
LOG = {"mukono": {"Housing": 153005, "Food": 713711, "NFNH": 664544, "Elder care (5%)": 76563, "Margins (5%)": 72741},
       "nakaseke": {"Housing": 211247, "Food": 643882, "NFNH": 649487, "Elder care (5%)": 75231, "Margins (5%)": 68430}}
LOG_V001 = {"mukono": {"Housing": 98297, "Food": 648048, "NFNH": 553121, "Elder care (5%)": 64973, "Margins (5%)": 63307},
            "nakaseke": {"Housing": 76891, "Food": 772623, "NFNH": 599498, "Elder care (5%)": 72451, "Margins (5%)": 72229}}
COMP = list(LOG["mukono"])


def rebuild_benchmark():
    """Rebuild the v0.02 workbooks with scripts/benchmark in a temp folder; return recalculated values or None."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        print("LibreOffice not found: using published benchmark values only"); return None
    import openpyxl
    sys.path.insert(0, str(BENCH)); import xl_compat
    tmp = Path(tempfile.mkdtemp(prefix="bm_v002_"))
    res = {}
    try:
        for d in DISTS:
            part, wbk = tmp / f"{d}_part.csv", tmp / f"{d}_v002.xlsx"
            src = RAW / f"NFC_LW Benchmark tool_Clean_v0.01 - {d.title()}.xlsx"
            kw = dict(check=True, cwd=BENCH, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            subprocess.run([sys.executable, "-W", "ignore", "build_participants.py", str(DER / "Survey_Cleaned_v1.csv"),
                            str(RAW / "Standardized_Data.csv"), str(part), d], **kw)
            subprocess.run([sys.executable, "-W", "ignore", "build_benchmark.py", d, str(src), str(part),
                            str(DER / "Market_Prices_Tidy_v1_1.csv"), str(wbk)], **kw)
            xl_compat.make_verify_copy(str(wbk), str(tmp / f"{d}_check.xlsx"))
        subprocess.run([soffice, f"-env:UserInstallation=file://{tmp}/lo_profile", "--headless", "--convert-to", "xlsx",
                        "--outdir", str(tmp / "calc")] + [str(tmp / f"{d}_check.xlsx") for d in DISTS],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=150)
        for d in DISTS:
            wb = openpyxl.load_workbook(tmp / "calc" / f"{d}_check.xlsx", data_only=True)
            o, f, b = wb["OUTPUT_Living wage"], wb["3. Food costs"], wb["Background calculations"]
            comp = dict(zip(COMP, [o[f"E{r}"].value for r in range(27, 32)]))
            if any(abs(comp[k] - LOG[d][k]) > 1 for k in COMP):
                print(f"WARNING {d}: rebuilt benchmark differs from the log {comp}; using published values only"); return None
            diet = pd.DataFrame([[f.cell(r, 3).value, f.cell(r, 4).value, f.cell(r, 5).value, f.cell(r, 6).value]
                                 for r in range(49, 68)], columns=["item", "edible_g_day", "purchase_g_day", "cost_ugx_day"])
            nf = pd.DataFrame([[b[f"B{r}"].value, b[f"C{r - 5}"].value, b[f"D{r - 5}"].value, b[f"C{r}"].value, b[f"G{r}"].value]
                               for r in (291, 292, 293)],
                              columns=["category", "cpi_weight_national", "tool_estimate_ugx_month", "survey_participants_ugx_month", "deviation_factor"])
            res[d] = dict(comp=comp, diet=diet, nfnh=nf,
                          scal=dict(diet_day=f["F68"].value, additions=sum(f[f"D{r}"].value for r in (71, 72, 73)),
                                    diet_day_total=f["D75"].value, adult_month=f["E21"].value, hh_food=f["E14"].value,
                                    energy_target=f["D35"].value, energy_calc=f["E35"].value,
                                    food_share=b["D275"].value, nfnh_share=b["D276"].value, nfnh_pre=b["C297"].value,
                                    nfnh_factor=b["C298"].value, nfnh_final=b["C300"].value),
                          prices=pd.read_csv(tmp / f"{d}_v002_food_price_log.csv"),
                          part=pd.read_csv(tmp / f"{d}_part.csv"))
            w1 = openpyxl.load_workbook(RAW / f"NFC_LW Benchmark tool_Clean_v0.01 - {d.title()}.xlsx", data_only=True)
            f1 = w1["3. Food costs"]
            res[d]["diet_v001"] = pd.DataFrame([[f1.cell(r, 3).value, f1.cell(r, 5).value, f1.cell(r, 6).value] for r in range(49, 68)],
                                               columns=["item", "purchase_g_day", "cost_ugx_day"])
            res[d]["scal_v001"] = dict(diet_day=f1["F68"].value, diet_day_total=f1["D75"].value, hh_food=f1["E14"].value)
        return res
    except Exception as e:                                # noqa: BLE001
        print(f"Benchmark rebuild failed ({e}); using published values only"); return None
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


BM = rebuild_benchmark()

# ---------------------------------------------------------------- costing
o = odir("costing")
rows = []
for d in DISTS:
    r = {"district": DIST_LABEL[d]}
    for k in COMP: r[f"{k} v0.02"] = LOG[d][k]
    r["Total v0.02 (UGX/month)"] = sum(LOG[d].values())
    r["Total v0.02 (UGX/year)"] = sum(LOG[d].values()) * 12
    r["Total v0.01 (UGX/month)"] = sum(LOG_V001[d].values())
    rows.append(r)
comp_tab = pd.DataFrame(rows)
comp_tab.to_csv(o / "benchmark_components_v002.csv", index=False)
to_tex(comp_tab[["district"] + [f"{k} v0.02" for k in COMP] + ["Total v0.02 (UGX/month)", "Total v0.01 (UGX/month)"]],
       o / "benchmark_components_v002.tex", "Living income benchmark v0.02 by component, reference household 2 adults + 3 children (UGX per month)",
       "tab:benchmark_components_v002")
if BM:
    items, summ = [], []
    for d in DISTS:
        x = BM[d]["diet"].copy(); x.insert(0, "district", DIST_LABEL[d])
        x["slot"] = range(1, len(x) + 1)
        x["price_ugx_kg"] = x.cost_ugx_day / x.purchase_g_day * 1000
        pl = BM[d]["prices"].drop_duplicates("tool_item").set_index("tool_item")
        x["price_basis"] = x["item"].map(pl.basis).fillna("tool value kept")
        x["market_prices_n"] = x["item"].map(pl.n_prices).fillna(0).astype(int)
        items.append(x[["district", "slot", "item", "edible_g_day", "purchase_g_day", "price_ugx_kg", "cost_ugx_day", "price_basis", "market_prices_n"]])
        s, s1 = BM[d]["scal"], BM[d]["scal_v001"]
        summ.append({"district": DIST_LABEL[d], "model diet, adult male (UGX/day)": s["diet_day"],
                     "+ salt, waste, variability (%)": s["additions"] * 100, "model diet incl. additions (UGX/day)": s["diet_day_total"],
                     "adult male (UGX/month, x30.43)": s["adult_month"], "household food cost v0.02 (UGX/month)": s["hh_food"],
                     "household food cost v0.01 (UGX/month)": s1["hh_food"], "energy target (kcal)": s["energy_target"],
                     "energy in diet (kcal)": s["energy_calc"],
                     "old notebook modal diet (UGX/month)": {"mukono": 167382, "nakaseke": 172128}[d]})
    items = pd.concat(items); items.round(1).to_csv(o / "model_diet_items_v002.csv", index=False)
    summ = pd.DataFrame(summ); summ.round(1).to_csv(o / "model_diet_cost_v002.csv", index=False)
    pd.concat([BM[d]["prices"].assign(district=DIST_LABEL[d]) for d in DISTS]).round(1).to_csv(o / "food_price_inputs_v002.csv", index=False)
    to_tex(summ[["district", "model diet, adult male (UGX/day)", "model diet incl. additions (UGX/day)", "adult male (UGX/month, x30.43)",
                 "household food cost v0.02 (UGX/month)", "household food cost v0.01 (UGX/month)"]],
           o / "model_diet_cost_v002.tex", "Cost of the benchmark model diet (tool v0.02, July 2025 district market prices)", "tab:model_diet_cost_v002")
    for d in DISTS:
        t = items[items.district == DIST_LABEL[d]][["item", "purchase_g_day", "price_ugx_kg", "cost_ugx_day", "price_basis"]]
        t.columns = ["Item", "g/day purchased", "UGX/kg", "UGX/day", "price source"]
        to_tex(t, o / f"model_diet_items_{d}_v002.tex", f"Model diet items and cost per adult male per day, {DIST_LABEL[d]} (tool v0.02)",
               f"tab:model_diet_items_{d}", fmt={"g/day purchased": "{:.0f}", "UGX/kg": "{:,.0f}", "UGX/day": "{:,.0f}"})

# ---------------------------------------------------------------- nfnh_identification
o = odir("nfnh_identification")
NFNH_INPUTS = pd.DataFrame([
    ("Education", "Education", "school fees + materials per term (by level) x children at that level x 3 terms / 12; university per semester x 2 / 12"),
    ("Healthcare", "Healthcare", "health spending in the last 3 months / 3 (column labelled 'number of visits' holds amounts)"),
    ("Transport", "Transport", "cost per trip x uses per month for each transport mode (everyday 30, twice 2, four times 4, more than 5 = 6)")],
    columns=["category", "participant_column", "survey rule (scripts/benchmark/build_participants.py)"])
if BM:
    nf = []
    for d in DISTS:
        x = BM[d]["nfnh"].copy(); x.insert(0, "district", DIST_LABEL[d])
        p = BM[d]["part"]
        x["participants_n"] = len(p)
        x["participant_hh_median_ugx_month"] = x.category.map({"Education": p.Education.median(), "Transport": p.Transport.median(), "Health": p.Healthcare.median()})
        s = BM[d]["scal"]
        x = pd.concat([x, pd.DataFrame([{"district": DIST_LABEL[d], "category": "NFNH before post-check (CPI share x food)",
                                         "tool_estimate_ugx_month": s["nfnh_pre"], "deviation_factor": s["nfnh_factor"]},
                                        {"district": DIST_LABEL[d], "category": "Final NFNH (published v0.02)",
                                         "tool_estimate_ugx_month": s["nfnh_final"]}])])
        nf.append(x)
    nf = pd.concat(nf); nf.round(3).to_csv(o / "nfnh_postcheck_v002.csv", index=False)
    t = nf[["district", "category", "cpi_weight_national", "tool_estimate_ugx_month", "survey_participants_ugx_month", "deviation_factor"]].copy()
    t.columns = ["District", "Category", "CPI weight", "tool (UGX/month)", "15 participants (UGX/month)", "deviation factor"]
    to_tex(t, o / "nfnh_postcheck_v002.tex", "Non-food non-housing (NFNH) costs in the v0.02 benchmark: CPI-share estimate vs participants' reported spending",
           "tab:nfnh_postcheck", fmt={"CPI weight": "{:.3f}", "deviation factor": "{:.3f}"})
NFNH_INPUTS.to_csv(o / "nfnh_survey_inputs.csv", index=False)

# ================================================================== 2. food comments (food_cleaning, food_constraints)
FG = {"cereals_1": "Cereals", "prepared_cereals": "Prepared cereals", "roots_and_tubers": "Roots & tubers",
      "starchy_fruits_1": "Starchy fruits (matooke)", "vegetables_1": "Vegetables", "fruits": "Fruits", "legumes_1": "Legumes",
      "diary_1": "Dairy", "eggs_1": "Eggs", "cooking_oil": "Cooking oil", "non_alcoholic": "Non-alcoholic drinks"}
# keyword codes exactly as module_d_living_conditions.py (D2)
CODES = [("Pests & diseases", r"pest|disease|worm|insect|mosaic|rot|blight|army|rat|bird|zibugo|kiwotokwa"),
         ("Price / cannot afford", r"expens|money|afford|price|cost"), ("Distance / market access", r"far|distance|town|market|transport|hawker|reach|shop"),
         ("Weather / drought", r"weather|sun|drought|rain|heat|dry"), ("Quality / safety", r"dilut|water|spoil|quality|fresh|cold|small"),
         ("Availability / seasonal", r"scarc|season|availab|hard to get|rare|once in a while|take long"),
         ("No constraint", r"^no |none|no problem|not .*problem|^nothing")]
CORDER = [k for k, _ in CODES] + ["Other"]
fcols = {c: c.split("group_wg2nh50_")[1].split("/")[0] for c in sv.columns if "group_wg2nh50_" in c and c.endswith("food_source_comments")}
rec = []
for c, g in fcols.items():
    s = sv[["_id", "District", c]].dropna()
    for _, r in s.iterrows():
        txt = str(r[c]).strip(); usable = len(txt) > 2
        hit = ([k for k, rx in CODES if re.search(rx, txt.lower())] or ["Other"]) if usable else []
        rec.append({"_id": r["_id"], "district": r["District"], "food_group": FG[g], "comment": txt, "usable": usable, "codes": "; ".join(hit)})
C = pd.DataFrame(rec)
o = odir("food_cleaning")
C.to_csv(o / "food_comments_long.csv", index=False)
cg = C[C.usable].groupby(["food_group", "district"])._id.nunique().unstack(fill_value=0).reindex(FG.values()).fillna(0).astype(int)
cg.columns = [f"{DIST_LABEL[k]} HH with comment" for k in cg.columns]
for d in DISTS: cg[f"{DIST_LABEL[d]} % of HH"] = (cg[f"{DIST_LABEL[d]} HH with comment"] / NHH[d] * 100).round(1)
anyc = C[C.usable].groupby("district")._id.nunique()
cg.loc["Any food group"] = [anyc.get("mukono", 0), anyc.get("nakaseke", 0), round(anyc.get("mukono", 0) / NHH["mukono"] * 100, 1),
                            round(anyc.get("nakaseke", 0) / NHH["nakaseke"] * 100, 1)]
cg.to_csv(o / "food_comments_by_group.csv")
pd.DataFrame({"n comments (non-empty)": [len(C)], "usable (>2 characters)": [int(C.usable.sum())],
              "dropped (blank/junk)": [int((~C.usable).sum())]}).to_csv(o / "food_comments_counts.csv", index=False)

# ---- food_constraints (Nakaseke shares)
o = odir("food_constraints")
L = C[C.usable].assign(constraint=lambda x: x.codes.str.split("; ")).explode("constraint")
hh = L.drop_duplicates(["_id", "constraint"]).groupby(["district", "constraint"])._id.nunique().unstack(0).reindex(CORDER).fillna(0)
ov = pd.DataFrame({f"{DIST_LABEL[d]} HH": hh.get(d, 0).astype(int) for d in DISTS})
for d in DISTS: ov[f"{DIST_LABEL[d]} % of all HH"] = (ov[f"{DIST_LABEL[d]} HH"] / NHH[d] * 100).round(1)
nk_any = int(anyc.get("nakaseke", 0))
ov["Nakaseke % of HH with a comment"] = (ov["Nakaseke HH"] / nk_any * 100).round(1)
ov.index.name = "constraint"; ov.to_csv(o / "overall_food_constraints.csv")
nkL = L[L.district == "nakaseke"]
ngrp = nkL.groupby("food_group")._id.nunique().reindex(FG.values())
mat = (nkL.drop_duplicates(["_id", "food_group", "constraint"]).groupby(["food_group", "constraint"])._id.nunique()
       .unstack(fill_value=0).reindex(index=FG.values(), columns=CORDER).fillna(0))
pct = (mat.div(ngrp, axis=0) * 100).round(1)
out = mat.astype(int).add_suffix(" (HH)").join(pct.add_suffix(" (%)")); out.insert(0, "HH commenting", ngrp.astype(int))
out.to_csv(o / "food_constraints_by_food_group_nakaseke.csv")
L[L.district == "mukono"].groupby(["food_group", "constraint"])._id.nunique().unstack(fill_value=0).to_csv(o / "food_constraints_by_food_group_mukono_counts.csv")

# fig 1: overall bar (Nakaseke)
sel = ov.sort_values("Nakaseke % of all HH")
fig, ax = plt.subplots(figsize=(6.6, 3.8)); y = np.arange(len(sel))
ax.barh(y, sel["Nakaseke % of all HH"], color=DIST["nakaseke"], height=0.65)
for i, v in enumerate(sel["Nakaseke % of all HH"]): ax.text(v + 0.8, i, f"{v:.0f}%", va="center", fontsize=8, color=INK2)
ax.set_yticks(y, sel.index); ax.grid(axis="y", visible=False); ax.set_xlim(0, max(60, sel["Nakaseke % of all HH"].max() + 8))
ax.set_xlabel("% of Nakaseke households mentioning the constraint (any food group)")
ax.set_title("Food access constraints named by Nakaseke households")
save(fig, o / "overall_food_constraints_bar.png",
     f"Nakaseke, n = {NHH['nakaseke']} HH ({nk_any} with at least one usable comment). Open comments per food group coded by keyword "
     "(as Table D2); one comment can carry several codes. Mukono omitted: enumerators recorded comments for few households "
     f"({int(anyc.get('mukono', 0))} of {NHH['mukono']}); counts in the CSV.")

# fig 2: food group x constraint matrix (merges old bubble scatter + stacked bar)
cm = LinearSegmentedColormap.from_list("blues", [SURF, ORD4[0], ORD4[1], ORD4[2], ORD4[3]])
fig, ax = plt.subplots(figsize=(8.2, 4.8))
im = ax.imshow(pct.values, cmap=cm, vmin=0, vmax=100, aspect="auto")
for i in range(pct.shape[0]):
    for j in range(pct.shape[1]):
        v = pct.values[i, j]
        if v > 0: ax.text(j, i, f"{v:.0f}", ha="center", va="center", fontsize=7.5, color=SURF if v >= 55 else INK)
ax.set_xticks(range(len(CORDER)), [textwrap.fill(c.replace(" / ", "/ "), 12, break_long_words=False) for c in CORDER], fontsize=7.5, rotation=0)
ax.set_yticks(range(len(pct)), [f"{g} (n={int(n)})" for g, n in zip(pct.index, ngrp.values)])
ax.grid(False); ax.tick_params(length=0)
for s in ax.spines.values(): s.set_visible(False)
cb = fig.colorbar(im, ax=ax, fraction=0.03, pad=0.02); cb.set_label("% of households commenting on the food group", color=INK2); cb.outline.set_visible(False)
ax.set_title("Constraints named for each food group, Nakaseke (% of households commenting on that group)")
save(fig, o / "food_constraints_by_food_group.png",
     "Nakaseke only; n = households with a usable comment for the food group. Keyword coding as Table D2; rows can exceed 100% "
     "(several codes per comment). 'Other' = no keyword matched.")

# ================================================================== 3. market prices (Food_Analysis, market_prices)
CAT = {"Cereals_and_Grains": "Cereals", "Prepared_Cereals": "Prepared cereals", "Roots_Tubers": "Roots & tubers",
       "Starchy_Fruits_Vegetables": "Starchy fruits & fruit", "Pulses_and_Legumes": "Legumes", "Leafy_Vegetables": "Leafy vegetables",
       "Other_Vegetables": "Other vegetables", "Meat": "Meat", "Fish": "Fish", "Dairy": "Dairy", "Eggs": "Eggs",
       "Oils_Fats_Non_Alcoholic_Beverages": "Oils, fats & drinks"}
ITEM = {"beaf": "beef", "cookingoil": "cooking oil", "ground_nuts": "groundnuts", "irish_potatoes": "Irish potatoes",
        "sweet_potatoes": "sweet potatoes", "dried_beans": "dried beans", "fresh_beans": "fresh beans", "dried_soybeans": "soybeans",
        "cow_peas": "cowpeas", "sweet_banana": "sweet banana", "banana": "matooke", "ebugga": "bbuga", "naan": "chapati",
        "nonalcoholicbeverage_1": "soft drink 1", "nonalcoholicbeverage_2": "soft drink 2", "oat_products": "oats", "jack_fruit": "jackfruit"}
M = mk[mk.category.isin(CAT) & mk.price_ugx.notna()].copy()
M["food_group"] = M.category.map(CAT); M["item_label"] = M["item"].map(lambda s: ITEM.get(s, s.replace("_", " ")))

# Food_Analysis: number of price observations by district and food group (old groups_by_district / group_summary)
o = odir("Food_Analysis")
gbd = (M.groupby(["district", "food_group"]).agg(n_price_observations=("price_ugx", "size"), n_items=("item", "nunique"),
                                                n_markets=("market", "nunique")).reset_index())
gbd["district"] = gbd.district.map(DIST_LABEL); gbd.to_csv(o / "groups_by_district.csv", index=False)
gs = M.groupby("food_group").agg(total_obs=("price_ugx", "size"), unique_items=("item", "nunique"), districts=("district", "nunique"),
                                 obs_outlier_flagged=("outlier_flag", "count"),
                                 obs_local_unit_guessed=("price_std_guess", "count")).reset_index().sort_values("total_obs", ascending=False)
gs.to_csv(o / "group_summary.csv", index=False)

# per-row comparable price: measured UGX/kg or /litre; else guessed local-unit conversion to kg/litre; else per piece/tray etc.
M = M[M.outlier_flag.isna()].copy()


def basis(r):
    if r.price_std_unit in ("kg", "litre") and pd.notna(r.price_std): return r.price_std, r.price_std_unit, "measured"
    if pd.notna(r.price_std_guess) and r.local_factor_unit in ("kg", "litre"): return r.price_std_guess, r.local_factor_unit, "local unit (guessed kg)"
    if pd.notna(r.price_std): return r.price_std, r.price_std_unit, "measured"
    return np.nan, None, None


M[["price", "unit", "basis"]] = M.apply(lambda r: pd.Series(basis(r)), axis=1)
M = M[M.price.notna()]
# one unit per item: kg/litre if the item has any, else its most common unit
prim = M.groupby("item").unit.agg(lambda s: next((u for u in ("kg", "litre") if (s == u).any()), s.mode().iloc[0]))
M = M[M.unit == M["item"].map(prim)].copy()
M["market_label"] = M.district.map(DIST_LABEL) + " - " + M.market + " (" + M.venue_type + ")"
o = odir("market_prices")
M[["district", "market", "venue_type", "survey_date", "food_group", "item", "item_label", "vendor_slot", "price_ugx", "price", "unit", "basis"]] \
    .to_csv(o / "market_prices_used.csv", index=False)
tab = []
for (fg, it, lab, u), x in M.groupby(["food_group", "item", "item_label", "unit"]):
    r = {"food_group": fg, "item": lab, "unit": f"UGX/{u.replace('per ', '')}", "guessed_unit_share_%": round((x.basis != "measured").mean() * 100)}
    for d in DISTS:
        y = x[x.district == d].price
        r.update({f"{DIST_LABEL[d]} n": len(y), f"{DIST_LABEL[d]} median": y.median(), f"{DIST_LABEL[d]} min": y.min(), f"{DIST_LABEL[d]} max": y.max(),
                  f"{DIST_LABEL[d]} CV %": (y.std() / y.mean() * 100) if len(y) >= 3 else np.nan})
    tab.append(r)
tab = pd.DataFrame(tab).sort_values(["food_group", "item"]); tab.round(1).to_csv(o / "district_item_prices.csv", index=False)
to_tex(tab[["food_group", "item", "unit", "Mukono n", "Mukono median", "Mukono min", "Mukono max", "Nakaseke n", "Nakaseke median", "Nakaseke min", "Nakaseke max"]],
       o / "district_item_prices.tex", "Food prices in the July 2025 market survey by district (outlier-flagged prices excluded; UGX per unit shown)",
       "tab:district_item_prices", fmt={c: "{:.0f}" for c in ["Mukono n", "Nakaseke n"]})
by_mkt = M.groupby(["food_group", "item_label", "unit", "market_label"]).price.median().unstack()
by_mkt.round(0).to_csv(o / "market_item_median_prices.csv")
MK = list(by_mkt.columns)   # sorted: Mukono Kasawo, Mukono Nakifuma, Nakaseke Kasangombe, Nakaseke Kiwoko
# relative price index: item median in market / item median over all observations x 100, then median over items in a group
idx = by_mkt.div(M.groupby(["food_group", "item_label", "unit"]).price.median(), axis=0) * 100
gidx = idx.groupby(level="food_group").median().reindex([v for v in CAT.values() if v in idx.index.get_level_values(0)])
gn = idx.notna().groupby(level="food_group").sum().reindex(gidx.index)
gidx.round(0).join(gn.add_suffix(" items")).to_csv(o / "market_food_group_price_index.csv")
t = gidx.round(0).reset_index(); t.columns = ["Food group"] + [c.replace(" - ", ": ") for c in MK]
to_tex(t, o / "market_food_group_price_index.tex", "Relative food prices by market and food group (median over items of item median in the market / overall item median x 100)",
       "tab:market_food_group_price_index", fmt={c: "{:.0f}" for c in t.columns[1:]})

# fig: item prices per kg/litre by market (merges combined / mukono / nakaseke market food group figures)
kg = M[M.unit.isin(["kg", "litre"])]
pm = kg.groupby(["food_group", "item_label", "unit", "market_label"]).price.median().unstack()
gord = [v for v in CAT.values() if v in pm.index.get_level_values(0)]
pm = pm.reindex(sorted(pm.index, key=lambda k: (gord.index(k[0]), -np.nanmedian(pm.loc[k].values))))
guessed = kg.groupby(["food_group", "item_label", "unit"]).basis.apply(lambda s: (s != "measured").any())
fig, ax = plt.subplots(figsize=(7.4, 9.2)); yy = np.arange(len(pm))[::-1]
sty = {MK[0]: ("o", DIST["mukono"], "none"), MK[1]: ("o", DIST["mukono"], DIST["mukono"]),
       MK[2]: ("s", DIST["nakaseke"], "none"), MK[3]: ("s", DIST["nakaseke"], DIST["nakaseke"])}
for i, (k, row) in zip(yy, pm.iterrows()):
    v = row.dropna()
    if len(v) > 1: ax.hlines(i, v.min(), v.max(), color=GRID, lw=2.5, zorder=1)
for m in MK:
    mk_, col, fc = sty[m]
    ax.scatter(pm[m], yy, marker=mk_, s=30, edgecolor=col, facecolor=fc, linewidth=1.3, zorder=3,
               label=m.replace("mukono", "Mukono").replace(" - ", ": "))
ax.set_yticks(yy, [f"{k[1]}{'*' if guessed.get(k, False) else ''} (/{k[2]})" for k in pm.index], fontsize=7.5)
ax.set_xscale("log"); ax.xaxis.set_major_formatter(plt.FuncFormatter(lambda x, p: f"{x:,.0f}"))
ax.set_xlabel("Median price, UGX per kg or per litre (log scale)"); ax.grid(axis="y", visible=False)
b = 0
for g in gord:
    n = sum(1 for k in pm.index if k[0] == g)
    ax.text(1.01, yy[b] - (n - 1) / 2, g, transform=ax.get_yaxis_transform(), fontsize=7.5, color=INK2, va="center")
    if b: ax.axhline(yy[b] + 0.5, color=GRID, lw=0.8)
    b += n
ax.legend(loc="lower center", bbox_to_anchor=(0.45, 1.0), ncol=2, fontsize=7.5, handletextpad=0.3)
ax.set_title("Food prices by market, July 2025 (median per market)", pad=34)
save(fig, o / "market_food_group_prices.png",
     f"{len(kg)} prices (vendors x visits) for {len(pm)} items sold by weight or volume; outlier-flagged prices excluded. "
     "Hollow = rural market, filled = urban. * includes local-unit prices converted with guessed kg factors (price_std_guess). "
     "Items sold per piece/tray are in district_item_prices.csv.")

# fig: price variability (CV across vendors and visits) by district
cv = tab[tab.unit.isin(["UGX/kg", "UGX/litre"])].dropna(subset=["Mukono CV %", "Nakaseke CV %"], how="all").copy()
cv = cv.assign(o=cv[["Mukono CV %", "Nakaseke CV %"]].max(axis=1)).sort_values("o")
fig, ax = plt.subplots(figsize=(6.4, 6.4)); yy = np.arange(len(cv))
ax.hlines(yy, cv[["Mukono CV %", "Nakaseke CV %"]].min(axis=1), cv[["Mukono CV %", "Nakaseke CV %"]].max(axis=1), color=GRID, lw=2.5)
for d, mkr in (("mukono", "o"), ("nakaseke", "s")):
    ax.scatter(cv[f"{DIST_LABEL[d]} CV %"], yy, marker=mkr, s=30, color=DIST[d], zorder=3, label=DIST_LABEL[d])
ax.set_yticks(yy, [f"{a} ({b})" for a, b in zip(cv["item"], cv.food_group)], fontsize=7.5); ax.grid(axis="y", visible=False)
ax.set_xlabel("Coefficient of variation of price across vendors and visits (%)"); ax.legend(loc="lower right")
ax.set_title("How much food prices vary within each district, July 2025")
save(fig, o / "food_group_price_variability.png",
     "CV = standard deviation / mean of UGX per kg or litre; shown when a district has >= 3 prices for the item. "
     "Each district: 2 markets x 2 visits, up to 3 vendors per item. Outlier-flagged prices excluded.")

# fig: heatmap relative price index by food group x market
fig, ax = plt.subplots(figsize=(6.6, 4.6))
dv = LinearSegmentedColormap.from_list("div", [SRC[2], SURF, SRC[3]])   # green = cheaper, amber = dearer (not district colours)
v = gidx.values.astype(float); lim = np.nanmax(np.abs(v - 100))
im = ax.imshow(v, cmap=dv, norm=TwoSlopeNorm(100, 100 - lim, 100 + lim), aspect="auto")
for i in range(v.shape[0]):
    for j in range(v.shape[1]):
        if not np.isnan(v[i, j]): ax.text(j, i, f"{v[i, j]:.0f}", ha="center", va="center", fontsize=8, color=INK)
        else: ax.text(j, i, "-", ha="center", va="center", fontsize=8, color=INK2)
ax.set_xticks(range(len(MK)), [m.replace(" - ", "\n").replace(" (", "\n(").replace("mukono", "Mukono") for m in MK], fontsize=8)
ax.set_yticks(range(len(gidx)), [f"{g} ({int(gn.loc[g].max())})" for g in gidx.index], fontsize=8)
ax.grid(False); ax.tick_params(length=0)
for s in ax.spines.values(): s.set_visible(False)
cb = fig.colorbar(im, ax=ax, fraction=0.035, pad=0.02); cb.set_label("price index (all-market item median = 100)", color=INK2); cb.outline.set_visible(False)
ax.set_title("Relative food prices by market and food group, July 2025")
save(fig, o / "heatmap_market_food_groups.png",
     "Index = item median in the market / item median over all four markets x 100, then the median over the items of the group "
     "(number of items in brackets). Each item in its own unit (kg, litre, piece, tray), so groups can be compared across units. "
     "Outlier-flagged prices excluded.")

# ================================================================== 4. living income gap, both districts
h["net_ratio_hh"] = h.net_central / h.bench_hh_year * 100
h["gap_hh_gross"] = h.bench_hh_year - h.gross; h["gap_hh_net"] = h.bench_hh_year - h.net_central
h["gap_ref_net"] = h.bench_ref_year - h.net_central
bins = [-np.inf, 25, 50, 100, np.inf]
h["li_group_net"] = pd.cut(h.net_ratio_hh, bins, labels=LI_GROUPS, right=False)
h["li_group"] = pd.Categorical(h.li_group, LI_GROUPS, ordered=True)


def summary(x):
    r = {"Households": len(x), "Benchmark, reference HH (UGX/yr)": x.bench_ref_year.iloc[0] if x.bench_ref_year.nunique() == 1 else np.nan,
         "Benchmark, reference HH (UGX/month)": x.bench_ref_year.iloc[0] / 12 if x.bench_ref_year.nunique() == 1 else np.nan,
         "Benchmark, HH-size specific, median (UGX/yr)": x.bench_hh_year.median()}
    for lab, c in [("Cash income", "income_total_mi"), ("Home-grown food (in-kind)", "inkind_food_central"), ("Tree products used at home", "tree_inkind"),
                   ("Gross income incl. in-kind", "gross"), ("Production costs", "cost_central"), ("Net income", "net_central")]:
        r[f"{lab}, mean (UGX/yr)"] = x[c].mean(); r[f"{lab}, median (UGX/yr)"] = x[c].median()
    r["Gross income, median (UGX/month)"] = x.gross.median() / 12; r["Net income, median (UGX/month)"] = x.net_central.median() / 12
    for lab, c in [("gross", "ratio_hh"), ("net", "net_ratio_hh")]: r[f"Income as % of HH benchmark, median ({lab})"] = x[c].median()
    r["% below reference benchmark (gross)"] = (x.gross < x.bench_ref_year).mean() * 100
    r["% below reference benchmark (net)"] = (x.net_central < x.bench_ref_year).mean() * 100
    r["% below HH benchmark (cash only)"] = (x.income_total_mi < x.bench_hh_year).mean() * 100
    r["% below HH benchmark (gross)"] = (x.gross < x.bench_hh_year).mean() * 100
    r["% below HH benchmark (net)"] = (x.net_central < x.bench_hh_year).mean() * 100
    r["Gap to HH benchmark, median (gross, UGX/yr)"] = x.gap_hh_gross.median(); r["Gap to HH benchmark, median (net, UGX/yr)"] = x.gap_hh_net.median()
    r["Gap to HH benchmark, mean (gross, UGX/yr)"] = x.gap_hh_gross.mean(); r["Gap to HH benchmark, mean (net, UGX/yr)"] = x.gap_hh_net.mean()
    r["Gap to reference benchmark, median (gross, UGX/yr)"] = x.gap_ref.median(); r["Gap to reference benchmark, median (net, UGX/yr)"] = x.gap_ref_net.median()
    r["Gap to HH benchmark, median (gross, UGX/month)"] = x.gap_hh_gross.median() / 12
    return pd.Series(r)


S = pd.DataFrame({DIST_LABEL[d]: summary(h[h.district == d]) for d in DISTS}); S["All"] = summary(h)
mk_, nk_ = h[h.district == "mukono"], h[h.district == "nakaseke"]
S["p (district)"] = np.nan
for lab, c in [("Income as % of HH benchmark, median (gross)", "ratio_hh"), ("Income as % of HH benchmark, median (net)", "net_ratio_hh"),
               ("Gross income incl. in-kind, median (UGX/yr)", "gross"), ("Net income, median (UGX/yr)", "net_central"),
               ("Cash income, median (UGX/yr)", "income_total_mi")]:
    S.loc[lab, "p (district)"] = stats.mannwhitneyu(mk_[c], nk_[c]).pvalue
for lab, cond in [("% below HH benchmark (gross)", h.gross < h.bench_hh_year), ("% below HH benchmark (net)", h.net_central < h.bench_hh_year),
                  ("% below reference benchmark (gross)", h.gross < h.bench_ref_year)]:
    S.loc[lab, "p (district)"] = stats.chi2_contingency(pd.crosstab(h.district, cond))[1]
o = odir("living_income_gap_mukono")
S.index.name = "statistic"; S.round(4).to_csv(o / "living_income_gap_summary.csv")
key = ["Households", "Benchmark, reference HH (UGX/yr)", "Benchmark, HH-size specific, median (UGX/yr)", "Cash income, median (UGX/yr)",
       "Gross income incl. in-kind, median (UGX/yr)", "Net income, median (UGX/yr)", "Income as % of HH benchmark, median (gross)",
       "Income as % of HH benchmark, median (net)", "% below HH benchmark (cash only)", "% below HH benchmark (gross)", "% below HH benchmark (net)",
       "Gap to HH benchmark, median (gross, UGX/yr)", "Gap to HH benchmark, median (net, UGX/yr)"]
t = S.loc[key].reset_index()
t["p (district)"] = t["p (district)"].map(lambda p: "" if pd.isna(p) else ("<0.001" if p < 0.001 else f"{p:.3f}"))
to_tex(t, o / "living_income_gap_summary.tex", "Living income gap by district: cash, gross (incl. home-grown food) and net income vs the v0.02 benchmark",
       "tab:living_income_gap_summary")

# ---- living_income_gap_mukono_outliers (figures; no outlier removal)
o = odir("living_income_gap_mukono_outliers")
# robustness of the old IQR rule (kept only as a check)
rob = []
for d in DISTS:
    x = h[h.district == d]
    q1, q3 = x.gross.quantile([0.25, 0.75]); up = q3 + 1.5 * (q3 - q1); y = x[x.gross <= up]
    for lab, z in (("all households (used)", x), (f"excluding IQR outliers (gross > {up / 1e6:.1f}M), old rule", y)):
        rob.append({"district": DIST_LABEL[d], "sample": lab, "n": len(z), "mean gross (UGX/yr)": z.gross.mean(), "median gross (UGX/yr)": z.gross.median(),
                    "% below HH benchmark (gross)": (z.gross < z.bench_hh_year).mean() * 100,
                    "median % of HH benchmark (gross)": z.ratio_hh.median()})
pd.DataFrame(rob).round(1).to_csv(o / "income_gap_outlier_sensitivity.csv", index=False)

# fig: ranked households vs benchmark (merges income_scatter_clean and _colored)
fig, axs = plt.subplots(1, 2, figsize=(9.2, 4.0), sharey=True)
CAP = 200
for ax, d in zip(axs, DISTS):
    x = h[h.district == d].sort_values("ratio_hh").reset_index(drop=True); rk = np.arange(1, len(x) + 1)
    ax.scatter(rk, x.net_ratio_hh.clip(-25, CAP), s=6, color=INK2, alpha=0.45, lw=0, label="net income")
    ax.scatter(rk, x.ratio_hh.clip(upper=CAP), s=8, color=DIST[d], lw=0, label="gross income incl. home-grown food")
    ax.axhline(100, color=INK, lw=1, ls="--"); ax.text(4, 103, "living income benchmark (household size)", fontsize=7.5, color=INK)
    bg, bn = (x.ratio_hh < 100).mean() * 100, (x.net_ratio_hh < 100).mean() * 100
    ax.text(0.03, 0.80, f"below benchmark:\ngross {bg:.0f}%   net {bn:.0f}%", transform=ax.transAxes, fontsize=8, color=INK)
    ax.set_title(f"{DIST_LABEL[d]} (n = {len(x)})"); ax.set_xlabel("Households ranked by gross income / benchmark")
    ax.set_ylim(-28, CAP + 5)
axs[0].set_ylabel("Income as % of household-size benchmark"); axs[1].legend(loc="upper left", bbox_to_anchor=(0.0, 0.75), markerscale=2)
fig.suptitle("Household income relative to the living income benchmark v0.02", x=0.0, ha="left", fontweight="bold", fontsize=10.5)
fig.tight_layout()
ncap = int((h.ratio_hh > CAP).sum() + (h.net_ratio_hh > CAP).sum()); nneg = int((h.net_ratio_hh < -25).sum())
save(fig, o / "income_scatter_clean_colored.png",
     f"All 597 HH (no outlier removal). Benchmark v0.02 scaled to each household's size (reference 2+3: Mukono 20.17M, Nakaseke 19.78M UGX/yr). "
     f"Cash income = mean of 20 imputations. {ncap} values above {CAP}% and {nneg} below -25% drawn at the axis limit.")

# fig: median income vs benchmark (old living_income_gap_comparison)
bars = [("Cash income", "income_total_mi", SRC[0]), ("Gross incl. home-grown food", "gross", SRC[2]), ("Net (gross - production costs)", "net_central", SRC[3])]
fig, ax = plt.subplots(figsize=(6.6, 3.9)); xx = np.arange(2); w = 0.2
for k, (lab, c, col) in enumerate(bars):
    v = [h.loc[h.district == d, c].median() for d in DISTS]
    ax.bar(xx + (k - 1.5) * w, v, w, color=col, label=f"median {lab.lower()}")
    for xi, vi in zip(xx + (k - 1.5) * w, v): ax.text(xi, vi + 2e5, f"{vi / 1e6:.1f}M", ha="center", fontsize=7.5, color=INK2)
bref = [h.loc[h.district == d, "bench_ref_year"].iloc[0] for d in DISTS]
bhh = [h.loc[h.district == d, "bench_hh_year"].median() for d in DISTS]
ax.bar(xx + 1.5 * w, bref, w, color=GRID, edgecolor=INK2, label="benchmark, reference HH (2 adults + 3 children)")
for xi, vi in zip(xx + 1.5 * w, bref): ax.text(xi, vi + 2e5, f"{vi / 1e6:.1f}M", ha="center", fontsize=7.5, color=INK2)
ax.set_xticks(xx, [DIST_LABEL[d] for d in DISTS]); ax.yaxis.set_major_formatter(plt.FuncFormatter(money)); ax.grid(axis="x", visible=False)
ax.set_ylabel("UGX per year"); ax.legend(loc="upper right", fontsize=7.5)
ax.set_ylim(0, max(bref) * 1.35)
ax.set_title("Median household income vs the living income benchmark v0.02")
save(fig, o / "living_income_gap_comparison.png",
     f"Mukono n = {NHH['mukono']}, Nakaseke n = {NHH['nakaseke']}. Medians over households (means in living_income_gap_summary.csv). "
     f"Median household-size benchmark: Mukono {bhh[0] / 1e6:.1f}M, Nakaseke {bhh[1] / 1e6:.1f}M. Production costs: net-income v1 central.")

# fig: LI groups, gross vs net (old living_income_stacked)
sh = []
for d in DISTS:
    for lab, c in (("gross", "li_group"), ("net", "li_group_net")):
        x = h.loc[h.district == d, c].value_counts(normalize=True).reindex(LI_GROUPS).fillna(0) * 100
        sh.append(pd.Series(x.values, index=LI_GROUPS, name=f"{DIST_LABEL[d]}\n{lab}"))
sh = pd.DataFrame(sh)
sh.round(1).rename(index=lambda s: s.replace("\n", " ")).to_csv(o / "living_income_groups.csv")
fig, ax = plt.subplots(figsize=(6.4, 3.4)); left = np.zeros(len(sh)); yy = np.arange(len(sh))[::-1]
for g, col in zip(LI_GROUPS, ORD4):
    ax.barh(yy, sh[g], left=left, color=col, height=0.62, label=f"{g} of benchmark", edgecolor=SURF, lw=0.6)
    for yi, l, v in zip(yy, left, sh[g]):
        if v >= 4: ax.text(l + v / 2, yi, f"{v:.0f}", ha="center", va="center", fontsize=7.5, color=INK if col == ORD4[0] else SURF)
    left += sh[g].values
ax.set_yticks(yy, sh.index); ax.set_xlim(0, 100); ax.grid(axis="y", visible=False)
ax.set_xlabel("% of households"); ax.legend(loc="lower center", bbox_to_anchor=(0.5, 1.0), ncol=4, fontsize=7.5)
ax.set_title("Households by income as % of the living income benchmark", pad=22)
save(fig, o / "living_income_stacked.png",
     f"Mukono n = {NHH['mukono']}, Nakaseke n = {NHH['nakaseke']}. Benchmark v0.02 scaled to household size; gross = cash + home-grown food + "
     "tree products used at home; net = gross - production costs (net v1 central). Groups as li_group in Household_Analysis_v1.")

# ================================================================== 5. gap drivers: descriptive only (model = paper1 table B3)
o = odir("gap_drivers_mukono")
shutil.copyfile(R / "paper1" / "tables" / "B3_drivers_regression.csv", o / "B3_drivers_regression.csv")
VARS = [("Household size", "hh_size", "c"), ("Land owned (acres)", "land_owned_ac", "c"), ("Coffee area (acres)", "coffee_ac", "c"),
        ("Coffee yield (kg FAQ/acre)", "coffee_yield_faq_ac", "c"), ("Kiboko price (UGX/kg)", "kiboko_price", "c"),
        ("Trees on farm (count)", "trees_on_farm", "c"), ("Number of cash income sources", "n_income_sources", "c"),
        ("Respondent age (years)", "resp_age", "c"), ("Respondent secondary education+ (%)", "resp_edu_secondary_plus", "b"),
        ("Female household head (%)", "female_head", "b"), ("Farmer group member (%)", "group_member", "b"),
        ("Extension contact (%)", "extension_contact", "b"), ("Credit access (%)", "credit_access", "b")]
G3 = ["<25%", "25-50%", ">=50%"]
h["li3"] = pd.Categorical(np.where(h.ratio_hh < 25, G3[0], np.where(h.ratio_hh < 50, G3[1], G3[2])), G3, ordered=True)
rows, cors = [], []
for d in DISTS + ["all"]:
    x = h if d == "all" else h[h.district == d]
    for lab, c, kind in VARS:
        s = num(x[c]); r = {"district": DIST_LABEL.get(d, "All"), "variable": lab}
        for g in G3:
            v = s[x.li3 == g].dropna()
            r[f"{g} (n)"] = len(v); r[f"{g}"] = v.mean() * 100 if kind == "b" else v.median()
        ok = s.notna()
        if kind == "b":
            tb = pd.crosstab(x.li3[ok], s[ok]); r["test"] = "chi-square"; r["p"] = stats.chi2_contingency(tb)[1] if tb.shape[1] == 2 else np.nan
        else:
            r["test"] = "Kruskal-Wallis"; r["p"] = stats.kruskal(*[s[(x.li3 == g) & ok] for g in G3]).pvalue
        rows.append(r)
        for lab2, rc in (("gross", "ratio_hh"), ("net", "net_ratio_hh")):
            rho, p = stats.spearmanr(s[ok], x.loc[ok, rc])
            cors.append({"district": DIST_LABEL.get(d, "All"), "variable": lab, "income": lab2, "n": int(ok.sum()), "spearman_rho": rho, "p": p})
desc = pd.DataFrame(rows); desc.round(4).to_csv(o / "gap_characteristics_by_li_group.csv", index=False)
cor = pd.DataFrame(cors); cor.round(4).to_csv(o / "gap_correlation_spearman.csv", index=False)
cg2 = cor[(cor.income == "gross") & (cor.district != "All")].pivot(index="variable", columns="district", values="spearman_rho")
cg2 = cg2.reindex([v[0] for v in VARS])[::-1]
cn = cor[(cor.income == "gross") & (cor.district != "All")].pivot(index="variable", columns="district", values="n").reindex(cg2.index)
fig, ax = plt.subplots(figsize=(6.6, 4.6)); yy = np.arange(len(cg2))
ax.hlines(yy, cg2.min(axis=1), cg2.max(axis=1), color=GRID, lw=2.5)
for d, m in (("mukono", "o"), ("nakaseke", "s")):
    ax.scatter(cg2[DIST_LABEL[d]], yy, marker=m, s=34, color=DIST[d], zorder=3, label=f"{DIST_LABEL[d]} (n = {NHH[d]})")
ax.axvline(0, color=INK2, lw=0.8)
lim = np.ceil(np.nanmax(np.abs(cg2.values)) * 10 + 0.5) / 10
ax.set_yticks(yy, cg2.index); ax.grid(axis="y", visible=False); ax.set_xlim(-lim, lim)
ax.set_xlabel("Spearman correlation with gross income as % of household-size benchmark"); ax.legend(loc="lower right")
ax.set_title("Household and farm characteristics associated with the income ratio, by district")
save(fig, o / "gap_drivers_spearman_by_district.png",
     "Descriptive, unadjusted rank correlations (both districts, gross income incl. home-grown food / household-size benchmark v0.02). "
     "n differs by variable (missing answers; Namyenya's tree counts excluded), see CSV. Income components are not used as 'drivers'. Adjusted associations: paper1 table B3 (copied here). Net-income correlations and "
     "p-values in gap_correlation_spearman.csv; group comparisons (Kruskal-Wallis / chi-square) in gap_characteristics_by_li_group.csv.")
print("done:", sorted(p.name for p in OUT.iterdir() if p.is_dir()))
